from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.dashboard.app import _dataframe_to_xlsx_bytes


@pytest.mark.parametrize(
    "value",
    [
        "=1+1",
        '=HYPERLINK("https://example.invalid", "open")',
        "+1+1",
        "-1+1",
        "@SUM(1,1)",
    ],
)
def test_formula_like_user_text_is_exported_as_text(value):
    data = pd.DataFrame([{"Комментарий": value}])

    workbook = load_workbook(BytesIO(_dataframe_to_xlsx_bytes(data, "Synthetic")), data_only=False)
    cell = workbook.active["A2"]

    assert cell.value == value
    assert cell.data_type == "s"
    assert cell.hyperlink is None


def test_formula_like_column_name_is_exported_as_text():
    data = pd.DataFrame([["safe"]], columns=["=1+1"])

    workbook = load_workbook(BytesIO(_dataframe_to_xlsx_bytes(data, "Synthetic")), data_only=False)

    assert workbook.active["A1"].value == "=1+1"
    assert workbook.active["A1"].data_type == "s"


def test_numeric_values_keep_numeric_cell_types():
    data = pd.DataFrame([{"integer": 42, "decimal": 12.5, "negative": -7.25, "numeric_text": "00123"}])

    workbook = load_workbook(BytesIO(_dataframe_to_xlsx_bytes(data, "Synthetic")), data_only=False)
    row = workbook.active[2]

    assert [(cell.value, cell.data_type) for cell in row] == [
        (42, "n"),
        (12.5, "n"),
        (-7.25, "n"),
        ("00123", "s"),
    ]
