"""Official monthly CPI loaders and purchasing-power calculations."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO, StringIO
import csv
import json
import re

from openpyxl import load_workbook
import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from src.data.sqlite_store import cpi_observations, save_cpi_observations


START_YEAR = 2015
MAX_CPI_WORKBOOK_BYTES = 2 * 1024 * 1024
_MONTHS_EN = {
    name: month for month, name in enumerate(
        ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"), 1)
}
_MONTHS_RU = {
    name: month for month, name in enumerate(
        ("январь", "февраль", "март", "апрель", "май", "июнь",
         "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"), 1)
}


class CPIUnavailableError(ValueError):
    """Raised when a real-value metric would require guessing a CPI observation."""


@dataclass(frozen=True)
class CPIRelease:
    currency: str
    observations: tuple[dict, ...]
    source_version: str
    payload_sha256: str
    published_on: str | None = None


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds").replace("+00:00", "Z")


def _retry_session() -> requests.Session:
    session = requests.Session()
    retries = Retry(
        total=2,
        read=2,
        connect=2,
        backoff_factor=0.8,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=["GET"],
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retries)
    session.mount("https://", adapter)
    session.headers["User-Agent"] = "FinRep/0.1 official-CPI-loader"
    return session


def _get(session: requests.Session, url: str, **kwargs) -> requests.Response:
    response = session.get(url, timeout=45, **kwargs)
    response.raise_for_status()
    return response


def _payload_hash(payload: bytes) -> str:
    return sha256(payload).hexdigest()


def _release_version(label: str, payload_hash: str) -> str:
    return f"{label}:{payload_hash[:16]}"


def _decimal(value, *, label: str) -> Decimal:
    try:
        parsed = Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, AttributeError) as exc:
        raise ValueError(f"invalid {label}: {value!r}") from exc
    if not parsed.is_finite() or parsed <= 0:
        raise ValueError(f"invalid {label}: {value!r}")
    return parsed


def _index_from_monthly_rates(rates: list[tuple[str, Decimal]]) -> tuple[dict, ...]:
    index = Decimal("100")
    observations = []
    for period, rate in sorted(rates):
        index = index * rate / Decimal("100")
        observations.append({"period": period, "index_value": format(index, "f")})
    return tuple(observations)


def fetch_bls_cpi(session: requests.Session, *, end_year: int | None = None) -> CPIRelease:
    end_year = end_year or date.today().year
    url = "https://api.bls.gov/publicAPI/v2/timeseries/data/CUUR0000SA0"
    payloads = []
    observations = []
    for start_year in range(START_YEAR, end_year + 1, 10):
        chunk_end = min(start_year + 9, end_year)
        response = _get(session, url, params={"startyear": start_year, "endyear": chunk_end})
        payloads.append(response.content)
        document = response.json()
        if document.get("status") != "REQUEST_SUCCEEDED":
            raise ValueError("BLS CPI request did not succeed")
        series = document.get("Results", {}).get("series", [])
        if len(series) != 1 or series[0].get("seriesID") != "CUUR0000SA0":
            raise ValueError("BLS CPI response has an unexpected series")
        for item in series[0].get("data", []):
            period_code = str(item.get("period", ""))
            if not re.fullmatch(r"M(0[1-9]|1[0-2])", period_code):
                continue
            try:
                value = _decimal(item.get("value"), label="BLS CPI value")
            except ValueError:
                continue
            observations.append({
                "period": f"{int(item['year']):04d}-{int(period_code[1:]):02d}",
                "index_value": format(value, "f"),
            })
    if not observations:
        raise ValueError("BLS CPI response has no monthly observations")
    digest = _payload_hash(b"\n".join(payloads))
    latest = max(item["period"] for item in observations)
    return CPIRelease("USD", tuple(sorted(observations, key=lambda item: item["period"])),
                      _release_version(latest, digest), digest)


def fetch_ons_cpi(session: requests.Session) -> CPIRelease:
    url = ("https://www.ons.gov.uk/generator?format=csv&uri=%2Feconomy%2F"
           "inflationandpriceindices%2Ftimeseries%2Fd7bt%2Fmm23")
    response = _get(session, url)
    payload = response.content
    rows = list(csv.reader(StringIO(response.text)))
    metadata = {row[0]: row[1] for row in rows if len(row) >= 2 and row[0]}
    if metadata.get("CDID") != "D7BT":
        raise ValueError("ONS CPI response has an unexpected series")
    published_on = datetime.strptime(metadata["Release date"], "%d-%m-%Y").date().isoformat()
    observations = []
    for row in rows:
        if len(row) < 2:
            continue
        match = re.fullmatch(r"(\d{4}) ([A-Z]{3})", row[0].strip())
        if not match or match.group(2) not in _MONTHS_EN:
            continue
        year = int(match.group(1))
        if year < START_YEAR:
            continue
        value = _decimal(row[1], label="ONS CPI value")
        observations.append({
            "period": f"{year:04d}-{_MONTHS_EN[match.group(2)]:02d}",
            "index_value": format(value, "f"),
        })
    if not observations:
        raise ValueError("ONS CPI response has no monthly observations")
    digest = _payload_hash(payload)
    return CPIRelease("GBP", tuple(observations),
                      _release_version(published_on, digest), digest, published_on)


def fetch_eurostat_cpi(session: requests.Session) -> CPIRelease:
    url = ("https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/"
           "prc_hicp_minr")
    response = _get(session, url, params={
        "lang": "en", "freq": "M", "unit": "I25", "coicop18": "TOTAL",
        "geo": "EA", "sinceTimePeriod": f"{START_YEAR}-01",
    })
    payload = response.content
    document = response.json()
    if document.get("id") != ["freq", "unit", "coicop18", "geo", "time"]:
        raise ValueError("Eurostat HICP response has unexpected dimensions")
    indexes = document.get("dimension", {}).get("time", {}).get("category", {}).get("index", {})
    values = document.get("value", {})
    observations = []
    for period, position in indexes.items():
        raw_value = values.get(str(position))
        if raw_value is None:
            continue
        value = _decimal(raw_value, label="Eurostat HICP value")
        observations.append({"period": period, "index_value": format(value, "f")})
    if not observations:
        raise ValueError("Eurostat HICP response has no monthly observations")
    updated = str(document.get("updated", ""))
    published_on = updated[:10] if re.fullmatch(r"\d{4}-\d{2}-\d{2}.*", updated) else None
    digest = _payload_hash(payload)
    return CPIRelease("EUR", tuple(observations),
                      _release_version(updated or max(indexes), digest), digest, published_on)


def _fetch_world_bank_gem_cpi(
    session: requests.Session, *, country: str, currency: str,
    end_year: int | None = None,
) -> CPIRelease:
    end_year = end_year or date.today().year
    url = f"https://api.worldbank.org/v2/country/{country}/indicator/CPTOTNSXN"
    response = _get(session, url, params={
        "source": 15,
        "date": f"{START_YEAR}M01:{end_year}M12",
        "format": "json",
        "per_page": 1000,
    })
    payload = response.content
    document = response.json()
    if not isinstance(document, list) or len(document) != 2:
        raise ValueError("World Bank GEM CPI response has an unexpected shape")
    metadata, rows = document
    if str(metadata.get("sourceid")) != "15" or not isinstance(rows, list):
        raise ValueError("World Bank GEM CPI response has an unexpected source")
    observations = []
    for item in rows:
        period_match = re.fullmatch(r"(\d{4})M(0[1-9]|1[0-2])", str(item.get("date", "")))
        if period_match is None or item.get("value") is None:
            continue
        indicator = item.get("indicator", {})
        if (indicator.get("id") != "CPTOTNSXN"
                or item.get("country", {}).get("id") != country):
            raise ValueError("World Bank GEM CPI response contains an unexpected series")
        value = _decimal(item["value"], label="World Bank GEM CPI value")
        observations.append({
            "period": f"{period_match.group(1)}-{period_match.group(2)}",
            "index_value": format(value, "f"),
        })
    if not observations:
        raise ValueError("World Bank GEM CPI response has no monthly observations")
    observations.sort(key=lambda item: item["period"])
    digest = _payload_hash(payload)
    latest = observations[-1]["period"]
    last_updated = str(metadata.get("lastupdated", ""))
    published_on = last_updated if re.fullmatch(r"\d{4}-\d{2}-\d{2}", last_updated) else None
    return CPIRelease(
        currency, tuple(observations),
        _release_version(published_on or latest, digest), digest, published_on,
    )


def fetch_world_bank_russia_cpi(
    session: requests.Session, *, end_year: int | None = None,
) -> CPIRelease:
    return _fetch_world_bank_gem_cpi(
        session, country="RUS", currency="RUB", end_year=end_year)


def _workbook_monthly_rates(payload: bytes, *, sheet_name: str | None,
                            first_year: int) -> list[tuple[str, Decimal]]:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    rates = []
    if sheet_name:
        rows = list(worksheet.iter_rows(values_only=True))
        year_columns = {
            column: int(value) for column, value in enumerate(rows[3])
            if isinstance(value, (int, float)) and int(value) >= first_year
        }
        for row in rows[5:17]:
            month = _MONTHS_RU.get(str(row[0]).strip().casefold())
            if month is None:
                continue
            for column, year in year_columns.items():
                if column >= len(row) or row[column] in (None, ""):
                    continue
                rates.append((f"{year:04d}-{month:02d}", _decimal(row[column], label="Rosstat CPI rate")))
        return rates

    current_year = None
    for row in worksheet.iter_rows(values_only=True):
        marker = row[0]
        if isinstance(marker, (int, float)) and int(marker) >= first_year:
            current_year = int(marker)
            continue
        month = _MONTHS_RU.get(str(marker).strip().casefold())
        if current_year is None or month is None or len(row) < 2 or row[1] in (None, ""):
            continue
        rates.append((f"{current_year:04d}-{month:02d}", _decimal(row[1], label="Kazakhstan CPI rate")))
    return rates


def fetch_kazakhstan_cpi(session: requests.Session) -> CPIRelease:
    page_url = "https://stat.gov.kz/ru/industries/economy/prices/dynamic-tables/?period=month"
    page = _get(session, page_url)
    soup = BeautifulSoup(page.text, "html.parser")
    title = soup.find(string=lambda value: value and
                      "Индекс потребительских цен и его составляющие" in value)
    row = title.find_parent(class_="divTableRow") if title else None
    link = row.find("a", href=True) if row else None
    if link is None:
        raise ValueError("Kazakhstan CPI workbook link was not found")
    url = requests.compat.urljoin(page.url, link["href"])
    response = _get(session, url)
    payload = response.content
    rates = _workbook_monthly_rates(payload, sheet_name=None, first_year=2022)
    if not rates:
        raise ValueError("Kazakhstan CPI workbook has no monthly observations")
    digest = _payload_hash(payload)
    latest = max(period for period, _rate in rates)
    published_match = re.search(r"\b(\d{2})\.(\d{2})\.(\d{4})\b", row.get_text(" ", strip=True))
    published_on = (
        f"{published_match.group(3)}-{published_match.group(2)}-{published_match.group(1)}"
        if published_match else None
    )
    return CPIRelease("KZT", _index_from_monthly_rates(rates),
                      _release_version(published_on or latest, digest), digest, published_on)


def fetch_kazakhstan_hybrid_cpi(
    session: requests.Session, *, end_year: int | None = None,
) -> CPIRelease:
    world_bank = _fetch_world_bank_gem_cpi(
        session, country="KAZ", currency="KZT", end_year=end_year)
    official = fetch_kazakhstan_cpi(session)
    world_bank_indexes = {
        item["period"]: _decimal(item["index_value"], label="World Bank GEM CPI value")
        for item in world_bank.observations
    }
    anchor = world_bank_indexes.get("2021-12")
    if anchor is None:
        raise ValueError("World Bank GEM CPI response has no Kazakhstan 2021-12 anchor")

    world_bank_source = {
        "provider_id": "world_bank_gem",
        "source_name": "World Bank Global Economic Monitor",
        "source_url": (
            "https://datacatalog.worldbank.org/search/dataset/0037798/"
            "global-economic-monitor"
        ),
    }
    official_source = {
        "provider_id": "stat_kz",
        "source_name": "Бюро национальной статистики Казахстана",
        "source_url": (
            "https://stat.gov.kz/ru/industries/economy/prices/"
            "dynamic-tables/?period=month"
        ),
    }
    observations = [
        {**item, **world_bank_source}
        for item in world_bank.observations
        if item["period"] <= "2021-12"
    ]
    for item in official.observations:
        if item["period"] < "2022-01":
            continue
        official_index = _decimal(item["index_value"], label="Kazakhstan CPI value")
        observations.append({
            "period": item["period"],
            "index_value": format(anchor * official_index / Decimal("100"), "f"),
            **official_source,
        })
    if not any(item["period"] >= "2022-01" for item in observations):
        raise ValueError("Kazakhstan CPI workbook has no observations after the splice point")
    observations.sort(key=lambda item: item["period"])
    digest = _payload_hash(
        f"{world_bank.payload_sha256}\n{official.payload_sha256}".encode())
    published_dates = [
        value for value in (world_bank.published_on, official.published_on) if value
    ]
    published_on = max(published_dates) if published_dates else None
    return CPIRelease(
        "KZT", tuple(observations),
        _release_version(published_on or observations[-1]["period"], digest),
        digest, published_on,
    )


def fetch_rosstat_cpi(session: requests.Session) -> CPIRelease:
    page_url = "https://rosstat.gov.ru/statistics/price?print=1"
    page = _get(session, page_url)
    match = re.search(
        r'href="([^"]*/storage/mediabank/ipc_mes_[^"]+\.xlsx)"', page.text,
        flags=re.IGNORECASE,
    )
    if match is None:
        raise ValueError("Rosstat CPI workbook link was not found")
    workbook_url = requests.compat.urljoin(page.url, match.group(1))
    response = _get(session, workbook_url)
    payload = response.content
    rates = _workbook_monthly_rates(payload, sheet_name="01", first_year=START_YEAR)
    if not rates:
        raise ValueError("Rosstat CPI workbook has no monthly observations")
    digest = _payload_hash(payload)
    release_label = workbook_url.rsplit("/", 1)[-1].removesuffix(".xlsx")
    return CPIRelease("RUB", _index_from_monthly_rates(rates),
                      _release_version(release_label, digest), digest)


_FETCHERS = {
    "RUB": fetch_world_bank_russia_cpi,
    "KZT": fetch_kazakhstan_hybrid_cpi,
    "USD": fetch_bls_cpi,
    "GBP": fetch_ons_cpi,
    "EUR": fetch_eurostat_cpi,
}


def refresh_official_cpi(database, *, currencies=None, session=None,
                         fetched_at: str | None = None) -> dict:
    """Refresh selected official series; successful providers are committed independently."""
    selected = [str(currency).upper() for currency in (currencies or _FETCHERS)]
    unknown = sorted(set(selected) - set(_FETCHERS))
    if unknown:
        raise ValueError(f"unsupported CPI currencies: {', '.join(unknown)}")
    session = session or _retry_session()
    fetched_at = fetched_at or _utc_now()
    results = []
    for currency in selected:
        try:
            release = _FETCHERS[currency](session)
            saved = save_cpi_observations(
                database,
                currency=currency,
                observations=list(release.observations),
                source_version=release.source_version,
                payload_sha256=release.payload_sha256,
                fetched_at=fetched_at,
                published_on=release.published_on,
            )
            results.append({"currency": currency, "status": "updated", **saved})
        except Exception as exc:
            results.append({"currency": currency, "status": "error", "message": str(exc)})
    return {
        "status": "done" if all(item["status"] == "updated" for item in results)
        else "partial" if any(item["status"] == "updated" for item in results)
        else "error",
        "results": results,
    }


def import_official_cpi_workbook(
    database,
    *,
    currency: str,
    payload: bytes,
    fetched_at: str | None = None,
    published_on: str | None = None,
) -> dict:
    """Import an official Rosstat or Kazakhstan CPI workbook as a manual fallback."""
    currency = currency.strip().upper()
    if currency not in {"RUB", "KZT"}:
        raise ValueError("official CPI workbook import is supported only for RUB and KZT")
    if not payload or len(payload) > MAX_CPI_WORKBOOK_BYTES:
        raise ValueError("CPI workbook must be a non-empty XLSX file up to 2 MB")
    rates = _workbook_monthly_rates(
        payload,
        sheet_name="01" if currency == "RUB" else None,
        first_year=START_YEAR if currency == "RUB" else 2022,
    )
    if not rates:
        raise ValueError("official CPI workbook has no recognized monthly observations")
    digest = _payload_hash(payload)
    source = (
        {
            "provider_id": "rosstat",
            "source_name": "Росстат",
            "source_url": "https://rosstat.gov.ru/statistics/price",
        }
        if currency == "RUB" else {
            "provider_id": "stat_kz",
            "source_name": "Бюро национальной статистики Казахстана",
            "source_url": (
                "https://stat.gov.kz/ru/industries/economy/prices/"
                "dynamic-tables/?period=month"
            ),
        }
    )
    saved = save_cpi_observations(
        database,
        currency=currency,
        observations=[{**item, **source} for item in _index_from_monthly_rates(rates)],
        source_version=_release_version(
            "rosstat-xlsx" if currency == "RUB" else "stat-kz-xlsx", digest),
        payload_sha256=digest,
        fetched_at=fetched_at or _utc_now(),
        published_on=published_on,
    )
    return {"status": "done", "results": [
        {"currency": currency, "status": "updated", **saved}
    ]}


def real_value(value, *, observation_period: str, base_period: str,
               indexes: dict[str, Decimal]) -> Decimal:
    """Express a nominal value from observation_period in prices of base_period."""
    missing = [period for period in (observation_period, base_period) if period not in indexes]
    if missing:
        raise CPIUnavailableError(
            f"Нет официального индекса инфляции за {', '.join(dict.fromkeys(missing))}. "
            "Реальная стоимость недоступна."
        )
    amount = Decimal(str(value))
    return amount * indexes[base_period] / indexes[observation_period]


def effective_cpi_indexes(database, currency: str) -> dict[str, Decimal]:
    return {
        row["period"]: row["index_value"]
        for row in cpi_observations(database, currency=currency)
    }
